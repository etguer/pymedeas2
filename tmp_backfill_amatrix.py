"""
Patch models/16sectors_qc/economy.xlsx:
Backfill A-matrix years 1995-2003 with 2004 values for both World and Quebec sheets.
Keeps a backup first.
"""
import shutil, zipfile, os, re
import openpyxl
import numpy as np
from lxml import etree

PATH   = 'models/16sectors_qc/economy.xlsx'
BACKUP = 'models/16sectors_qc/economy_pre_backfill_backup.xlsx'
NS     = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'

shutil.copy(PATH, BACKUP)
print(f'Backed up to {BACKUP}')

def parse_ref(t):
    t = t.replace("'", "")
    m = re.match(r'([^!]+)!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?', t)
    if not m: return None
    def cn(c):
        n = 0
        for ch in c: n = n * 26 + (ord(ch) - 64)
        return n
    c1, r1 = cn(m.group(2)), int(m.group(3))
    c2 = cn(m.group(4)) if m.group(4) else c1
    r2 = int(m.group(5)) if m.group(5) else r1
    return m.group(1), r1, c1, r2, c2

def get_defined_names(path):
    """Returns list of (name, sheet, ref_text) to preserve duplicates across sheets."""
    with zipfile.ZipFile(path) as z:
        tree = etree.fromstring(z.read('xl/workbook.xml'))
        sheets = [s.get('name') for s in tree.findall(f'.//{{{NS}}}sheet')]
        entries = []
        for dn in tree.findall(f'.//{{{NS}}}definedName'):
            sid = dn.get('localSheetId')
            sheet = sheets[int(sid)] if sid else None
            entries.append((dn.get('name'), sheet, dn.text or ''))
    return entries

wb = openpyxl.load_workbook(PATH)
entries = get_defined_names(PATH)

# Build lookup: (name, sheet) -> ref_text
lookup = {(n, s): r for n, s, r in entries}

patched = 0
for name, sheet_name, ref_text in entries:
    if not name.startswith('historic_A_Matrix_year') or sheet_name not in wb.sheetnames:
        continue
    year_str = name.replace('historic_A_Matrix_year', '')
    try:
        year = int(year_str)
    except ValueError:
        continue
    if year > 2003:
        continue   # only backfill 1995-2003

    ref = parse_ref(ref_text)
    if not ref: continue
    _, r1, c1, r2, c2 = ref

    # Read 2004 A-matrix values from the same sheet
    ref2004_text = lookup.get(('historic_A_Matrix_year2004', sheet_name))
    if not ref2004_text:
        print(f'  SKIP {name} [{sheet_name}]: no 2004 reference')
        continue
    ref2004 = parse_ref(ref2004_text)
    if not ref2004: continue
    _, r1_04, c1_04, r2_04, c2_04 = ref2004

    ws = wb[sheet_name]

    # Read 2004 values
    src_rows = list(ws.iter_rows(min_row=r1_04, max_row=r2_04,
                                  min_col=c1_04, max_col=c2_04,
                                  values_only=True))

    # Check dimensions match
    rows_src = r2_04 - r1_04 + 1
    cols_src = c2_04 - c1_04 + 1
    rows_dst = r2 - r1 + 1
    cols_dst = c2 - c1 + 1

    if rows_src != rows_dst or cols_src != cols_dst:
        print(f'  SKIP {name}: size mismatch {rows_src}x{cols_src} vs {rows_dst}x{cols_dst}')
        continue

    # Write 2004 values into this year's range
    non_nan = sum(1 for row in src_rows for v in row if v is not None and v == v)
    for ri, row in enumerate(src_rows):
        for ci, val in enumerate(row):
            ws.cell(row=r1 + ri, column=c1 + ci).value = val

    print(f'  Backfilled [{sheet_name}] {name} with 2004 values ({non_nan} non-NaN cells)')
    patched += 1

wb.save(PATH)
print(f'\nDone. Backfilled {patched} A-matrix year ranges.')
