"""
Copy HFC/PFC/SF6 emission named ranges from 14-sector climate.xlsx
into 16-sector climate.xlsx. These are global parameters, not sector-specific.
Keeps a backup of the 16-sector file first.
"""
import zipfile, shutil, os, re
import openpyxl
from lxml import etree

SRC    = 'models/14sectors_cat/climate.xlsx'
DST    = 'models/16sectors_qc/climate.xlsx'
BACKUP = 'models/16sectors_qc/climate_pre_emissions_backup.xlsx'
NS     = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'

# Named ranges to copy (all in World sheet, same structure in both files)
RANGES_TO_COPY = [
    'HFC125_emissions', 'HFC134a_emissions', 'HFC143a_emissions',
    'HFC152a_emissions', 'HFC227ea_emissions', 'HFC23_emissions',
    'HFC245ca_emissions', 'HFC32_emissions', 'HFC4310mee_emissions',
    'PFCs_emissions', 'SF6_emissions',
]

shutil.copy(DST, BACKUP)
print(f'Backed up to {BACKUP}')

def parse_ref(t):
    t = t.replace("'", "")
    m = re.match(r'([^!]+)!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?', t)
    if not m: return None
    def cn(c):
        n = 0
        for ch in c: n = n * 26 + (ord(ch) - 64)
        return n
    c1, r1 = cn(m.group(2)), int(m.group(3))
    c2 = cn(m.group(4)) if m.group(4) else c1
    r2 = int(m.group(5)) if m.group(5) else r1
    return m.group(1), r1, c1, r2, c2

def get_defined_names(path):
    with zipfile.ZipFile(path) as z:
        tree = etree.fromstring(z.read('xl/workbook.xml'))
        sheets = [s.get('name') for s in tree.findall(f'.//{{{NS}}}sheet')]
        names = {}
        for dn in tree.findall(f'.//{{{NS}}}definedName'):
            sid = dn.get('localSheetId')
            names[dn.get('name')] = (sheets[int(sid)] if sid else None, dn.text or '')
    return names

# Read source values
n14 = get_defined_names(SRC)
wb14 = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

# Load destination for writing
wb16 = openpyxl.load_workbook(DST)
n16  = get_defined_names(DST)

copied = 0
for name in RANGES_TO_COPY:
    if name not in n14 or name not in n16:
        print(f'  SKIP {name} — not in both files')
        continue

    ref14 = parse_ref(n14[name][1])
    ref16 = parse_ref(n16[name][1])
    if not ref14 or not ref16:
        print(f'  SKIP {name} — could not parse ref')
        continue

    sheet14, r1_s, c1_s, r2_s, c2_s = ref14
    sheet16, r1_d, c1_d, r2_d, c2_d = ref16

    ws14 = wb14[sheet14]
    ws16 = wb16[sheet16]

    # Read all values from source
    src_rows = list(ws14.iter_rows(min_row=r1_s, max_row=r2_s,
                                    min_col=c1_s, max_col=c2_s,
                                    values_only=True))

    # Check dimensions match
    rows14 = r2_s - r1_s + 1
    cols14 = c2_s - c1_s + 1
    rows16 = r2_d - r1_d + 1
    cols16 = c2_d - c1_d + 1

    if rows14 != rows16 or cols14 != cols16:
        print(f'  SKIP {name} — size mismatch: 14sec {rows14}x{cols14} vs 16sec {rows16}x{cols16}')
        continue

    # Write values to destination
    for ri, row in enumerate(src_rows):
        for ci, val in enumerate(row):
            ws16.cell(row=r1_d + ri, column=c1_d + ci).value = val

    non_zero = sum(1 for row in src_rows for v in row
                   if v is not None and v != 'na'
                   and not (isinstance(v, float) and v == 0.0))
    print(f'  Copied {name}: {rows14}x{cols14} cells ({non_zero} non-zero values)')
    copied += 1

wb16.save(DST)
print(f'\nDone. Copied {copied} named ranges into {DST}')
