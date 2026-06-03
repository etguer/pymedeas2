"""
Fast xlsx comparison using iter_rows (bulk read) instead of cell-by-cell.
"""
import zipfile, re, sys
import openpyxl
import numpy as np
from lxml import etree

NS    = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
DIR14 = 'models/14sectors_cat'
DIR16 = 'models/16sectors_qc'
XLSX_FILES = ['climate.xlsx', 'economy.xlsx', 'energy.xlsx', 'land.xlsx',
              'materials.xlsx', 'parameters.xlsx', 'transport.xlsx', 'water.xlsx']
MAX_PARTIAL = 1000

def get_defined_names(path):
    with zipfile.ZipFile(path) as z:
        tree = etree.fromstring(z.read('xl/workbook.xml'))
        sheets = [s.get('name') for s in tree.findall(f'.//{{{NS}}}sheet')]
        names = {}
        for dn in tree.findall(f'.//{{{NS}}}definedName'):
            sid = dn.get('localSheetId')
            names[dn.get('name')] = (sheets[int(sid)] if sid else None, dn.text or '')
    return names

def parse_ref(t):
    t = t.replace("'", "")
    m = re.match(r'([^!]+)!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?', t)
    if not m: return None
    def cn(c):
        n=0
        for ch in c: n=n*26+(ord(ch)-64)
        return n
    c1,r1 = cn(m.group(2)), int(m.group(3))
    c2 = cn(m.group(4)) if m.group(4) else c1
    r2 = int(m.group(5)) if m.group(5) else r1
    return m.group(1), r1, c1, r2, c2

def has_val(v):
    if v is None or v == 'na': return False
    if isinstance(v, str): return v.startswith('=') or bool(v.strip())
    try:
        f = float(v)
        return not (np.isnan(f) or f == 0.0)
    except: return True

def bulk_read(wb, sheet, r1, c1, r2, c2):
    if sheet not in wb.sheetnames: return None
    ws = wb[sheet]
    vals = []
    for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True):
        vals.extend(row)
    return vals

total = 0
for fname in XLSX_FILES:
    print(f'Processing {fname}...', flush=True)
    p14, p16 = f'{DIR14}/{fname}', f'{DIR16}/{fname}'
    try:
        n14 = get_defined_names(p14)
        n16 = get_defined_names(p16)
        wb14 = openpyxl.load_workbook(p14, read_only=True, data_only=False)
        wb16 = openpyxl.load_workbook(p16, read_only=True, data_only=False)
    except Exception as e:
        print(f'  ERROR: {e}', flush=True); continue

    file_hdr = False
    for name in sorted(set(n14) & set(n16)):
        r14 = parse_ref(n14[name][1])
        r16 = parse_ref(n16[name][1])
        if not r14 or not r16: continue
        v14 = bulk_read(wb14, r14[0], r14[1], r14[2], r14[3], r14[4])
        v16 = bulk_read(wb16, r16[0], r16[1], r16[2], r16[3], r16[4])
        if v14 is None or v16 is None: continue

        d14 = [v for v in v14 if has_val(v)]
        d16 = [v for v in v16 if has_val(v)]

        msg = None
        if d14 and not d16:
            msg = f'  [ALL EMPTY] {name}  (14sec: {len(d14)}/{len(v14)} non-zero, 16sec: 0/{len(v16)})'
        elif d14 and len(v14)==len(v16) and len(v14)<=MAX_PARTIAL:
            gaps = [i for i,(a,b) in enumerate(zip(v14,v16)) if has_val(a) and not has_val(b)]
            if gaps:
                msg = f'  [PARTIAL]   {name}  ({len(gaps)}/{len(v14)} cells: 14sec has data, 16sec is 0/NA)'

        if msg:
            if not file_hdr:
                print(f'\n=== {fname} ===', flush=True); file_hdr=True
            print(msg, flush=True)
            total += 1

    print(f'  done.', flush=True)

print(f'\nTotal issues: {total}', flush=True)
