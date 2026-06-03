import openpyxl
wb = openpyxl.load_workbook('models/16sectors_qc/economy.xlsx', read_only=True)
print('Sheets:', wb.sheetnames)
ws = wb['Global']
print('Global sheet contents:')
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if any(v is not None for v in row):
        print(f'  row {i+1}: {row}')
