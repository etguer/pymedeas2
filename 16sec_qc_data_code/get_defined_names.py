
import openpyxl
import os.path

os.chdir('C:/Users/thinkpad/PycharmProjects/pymedeas2qc/')
wb = openpyxl.load_workbook("./pymedeas2/16sec_qc_data_code/original_and_mock_xlsx/14sec_original_data/economy.xlsx")

ws = wb['World']
for key in ws.defined_names:
    print(key)

ws = wb['Europe']
for key in ws.defined_names:
    print(key)

wb = openpyxl.load_workbook("./pymedeas2/16sec_qc_data_code/original_and_mock_xlsx/14sec_original_data/energy.xlsx")
ws_original = wb['World']
print(ws.defined_names.keys())
for key in ws.defined_names:
    print(key)

wb = openpyxl.load_workbook("./pymedeas2/16sec_qc_data_code/original_and_mock_xlsx/16sec_mock_data/energy.xlsx")
ws_mock = wb['World']
for key in ws.defined_names:
    print(key)



wb = openpyxl.load_workbook("./pymedeas2/16sec_qc_data_code/original_and_mock_xlsx/14sec_original_data/energy.xlsx")
ws_original = wb['Europe']
for key in ws.defined_names:
    print(key)




wb = openpyxl.load_workbook("./pymedeas2/16sec_qc_data_code/original_and_mock_xlsx/14sec_original_data/energy.xlsx")
ws_original = wb['World']
wb = openpyxl.load_workbook("./pymedeas2/16sec_qc_data_code/original_and_mock_xlsx/14sec_original_data/updated/energy.xlsx")
ws_original_updated = wb['World']
wb = openpyxl.load_workbook("./pymedeas2/16sec_qc_data_code/original_and_mock_xlsx/16sec_mock_data/energy.xlsx")
ws_mock = wb['World']

for key in ws_original_updated.defined_names:
    if key not in ws_original.defined_names:
        print(ws_original_updated.defined_names[key])


for key in ws_original.defined_names:
    if key not in ws_mock.defined_names:
        print(ws_original.defined_names[key])

len(ws_mock.defined_names)
len(ws_original.defined_names)
len(ws_original_updated.defined_names)